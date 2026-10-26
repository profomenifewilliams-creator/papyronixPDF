"""Template-based business document creation. No language model or network calls."""
from __future__ import annotations

import csv
import io
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class DocumentBuilderError(ValueError):
    """A safe, user-facing document-builder validation error."""


DOCUMENT_TYPES = {
    "business-letter": "Business letter",
    "project-report": "Project or business report",
    "meeting-minutes": "Meeting minutes",
    "resume": "CV / resume",
    "invoice": "Invoice",
    "proposal": "Project proposal",
}

MAX_TEXT_FIELD = 20000


def _text(values: dict[str, str], key: str, default: str = "") -> str:
    return (values.get(key) or default).strip()


def _require(values: dict[str, str], keys: list[str], message: str) -> None:
    if any(not _text(values, key) for key in keys):
        raise DocumentBuilderError(message)


def _paragraphs(value: str) -> list[str]:
    return [block.strip() for block in re.split(r"\n\s*\n", value or "") if block.strip()]


def _lines(value: str) -> list[str]:
    return [line.strip().lstrip("•*- ").strip() for line in (value or "").splitlines() if line.strip()]


def _add_body(doc: Document, value: str) -> None:
    for block in _paragraphs(value):
        doc.add_paragraph(block)


def _add_bullets(doc: Document, value: str, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for item in _lines(value):
        doc.add_paragraph(item, style=style)


def _add_section(doc: Document, title: str, value: str, bullets: bool = False) -> None:
    if not value.strip():
        return
    doc.add_heading(title, level=1)
    if bullets:
        _add_bullets(doc, value)
    else:
        _add_body(doc, value)


def _set_cell_shading(cell, fill: str = "EAF1FB") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])


def _new_document(title: str, author: str, organization: str, date_text: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(48, 65, 87)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor(37, 91, 174)
    doc.styles["Title"].font.size = Pt(27)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(11.5)
    header = section.header.paragraphs[0]
    header.text = "PAPYRONIX  |  DOCUMENT BUILDER"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor(91, 111, 139)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated locally with Papyronix  •  Review before issuing  •  Page ").font.size = Pt(8)
    _add_page_field(footer)
    if title:
        doc.core_properties.title = title[:250]
    doc.core_properties.author = author[:200] or "Papyronix Document Builder"
    doc.core_properties.subject = "Template-generated document"
    return doc


def _top_identity(doc: Document, values: dict[str, str], title: str, date_key: str = "document_date") -> None:
    doc.add_heading(title, level=0)
    metadata = []
    for key in ("full_name", "organization", "email", "phone", "address"):
        value = _text(values, key)
        if value:
            metadata.append(value)
    if metadata:
        p = doc.add_paragraph("  |  ".join(metadata))
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(95, 112, 136)
    date_text = _text(values, date_key)
    if date_text:
        p = doc.add_paragraph(f"Date: {date_text}")
        if p.runs:
            p.runs[0].font.size = Pt(9)


def _build_letter(doc: Document, v: dict[str, str]) -> None:
    _require(v, ["recipient_name", "subject", "body"], "For a letter, enter the recipient, subject and letter body.")
    if not (_text(v, "signatory_name") or _text(v, "full_name")):
        raise DocumentBuilderError("Enter your name in the Sender name field so the letter can be signed.")
    _top_identity(doc, v, _text(v, "document_title", "Business Letter"))
    recipient_lines = [_text(v, "recipient_name"), _text(v, "recipient_organization"), _text(v, "recipient_address")]
    for line in recipient_lines:
        if line:
            doc.add_paragraph(line)
    doc.add_paragraph()
    subject = doc.add_paragraph()
    subject.add_run("Subject: ").bold = True
    subject.add_run(_text(v, "subject"))
    greeting = _text(v, "greeting", f"Dear {_text(v, 'recipient_name')}")
    doc.add_paragraph(greeting if greeting.endswith((",", ":")) else greeting + ",")
    _add_body(doc, _text(v, "body"))
    closing = _text(v, "closing", "Yours sincerely,")
    doc.add_paragraph(closing)
    doc.add_paragraph(_text(v, "signatory_name", _text(v, "full_name")))
    if _text(v, "signatory_role"):
        doc.add_paragraph(_text(v, "signatory_role"))


def _build_report(doc: Document, v: dict[str, str]) -> None:
    title = _text(v, "document_title", "Project Report")
    if not any(_text(v, key) for key in ("executive_summary", "background", "findings", "recommendations", "conclusion")):
        raise DocumentBuilderError("Enter at least one report section, such as summary, findings or recommendations.")
    _top_identity(doc, v, title)
    _add_section(doc, "Executive summary", _text(v, "executive_summary"))
    _add_section(doc, "Background and objectives", _text(v, "background"))
    _add_section(doc, "Findings", _text(v, "findings"), bullets=True)
    _add_section(doc, "Recommendations", _text(v, "recommendations"), bullets=True)
    _add_section(doc, "Conclusion", _text(v, "conclusion"))


def _build_minutes(doc: Document, v: dict[str, str]) -> None:
    _require(v, ["meeting_title", "meeting_date"], "For meeting minutes, enter the meeting title and date.")
    _top_identity(doc, v, _text(v, "meeting_title"), "meeting_date")
    details = [("Location", "meeting_location"), ("Chair / facilitator", "chair")]
    for label, key in details:
        if _text(v, key):
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(_text(v, key))
    _add_section(doc, "Attendees", _text(v, "attendees"), bullets=True)
    _add_section(doc, "Agenda", _text(v, "agenda"), bullets=True)
    _add_section(doc, "Discussion notes", _text(v, "discussion"))
    _add_section(doc, "Decisions", _text(v, "decisions"), bullets=True)
    actions = _lines(_text(v, "action_items"))
    if actions:
        doc.add_heading("Action items", level=1)
        parsed = []
        for action in actions:
            cells = [part.strip() for part in action.split("|")]
            if len(cells) >= 3:
                parsed.append(cells[:3])
            else:
                parsed.append([action, "Not assigned", "Not set"])
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ("Action", "Owner", "Due date")
        for cell, label in zip(table.rows[0].cells, headers):
            cell.text = label
            _set_cell_shading(cell, "DCE8FA")
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for action in parsed:
            cells = table.add_row().cells
            for cell, value in zip(cells, action):
                cell.text = value


def _build_resume(doc: Document, v: dict[str, str]) -> None:
    _require(v, ["full_name"], "Enter the person's name to create a CV/resume.")
    if not any(_text(v, key) for key in ("professional_summary", "skills", "experience", "education")):
        raise DocumentBuilderError("Add at least a professional summary, skills, work experience or education.")
    p = doc.add_heading(_text(v, "full_name"), level=0)
    if _text(v, "headline"):
        p = doc.add_paragraph(_text(v, "headline"))
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(37, 91, 174)
    contact = [x for x in (_text(v, "email"), _text(v, "phone"), _text(v, "address")) if x]
    if contact:
        p = doc.add_paragraph("  |  ".join(contact))
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(95, 112, 136)
    _add_section(doc, "Professional profile", _text(v, "professional_summary"))
    _add_section(doc, "Key skills", _text(v, "skills"), bullets=True)
    _add_section(doc, "Work experience", _text(v, "experience"))
    _add_section(doc, "Education", _text(v, "education"))
    _add_section(doc, "Certifications and training", _text(v, "certifications"), bullets=True)
    _add_section(doc, "References", _text(v, "references"))


def _parse_invoice_items(value: str) -> list[tuple[str, Decimal, Decimal]]:
    items = []
    for row_number, line in enumerate(_lines(value), start=1):
        parts = [part.strip() for part in next(csv.reader([line], delimiter="|"))]
        if len(parts) != 3:
            raise DocumentBuilderError(f"Invoice line {row_number} must use: Description | Quantity | Unit price.")
        description, qty_raw, price_raw = parts
        try:
            qty = Decimal(qty_raw.replace(",", ""))
            price = Decimal(price_raw.replace(",", "").replace("₦", "").replace("$", "").replace("€", ""))
        except InvalidOperation as exc:
            raise DocumentBuilderError(f"Invoice line {row_number} has an invalid quantity or unit price.") from exc
        if not description or qty <= 0 or price < 0 or not qty.is_finite() or not price.is_finite():
            raise DocumentBuilderError(f"Invoice line {row_number} needs a description, positive quantity and valid non-negative price.")
        if len(description) > 250:
            raise DocumentBuilderError(f"Invoice line {row_number} description must be 250 characters or fewer.")
        items.append((description, qty, price))
    if not items:
        raise DocumentBuilderError("Add at least one invoice line item.")
    if len(items) > 100:
        raise DocumentBuilderError("Invoices support up to 100 line items.")
    return items


def _fmt_decimal(value: Decimal) -> str:
    return f"{value:,.2f}"


def _build_invoice(doc: Document, v: dict[str, str]) -> None:
    _require(v, ["full_name", "recipient_name", "invoice_number", "line_items"], "For an invoice, enter your business/sender name, client, invoice number and line items.")
    items = _parse_invoice_items(_text(v, "line_items"))
    currency = _text(v, "currency", "NGN")[:12]
    _top_identity(doc, v, "INVOICE", "invoice_date")
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for label, value in (("Invoice number", _text(v, "invoice_number")), ("Issue date", _text(v, "invoice_date")), ("Due date", _text(v, "due_date"))):
        if value:
            cells = info.add_row().cells
            cells[0].text, cells[1].text = label, value
            _set_cell_shading(cells[0], "EAF1FB")
            for run in cells[0].paragraphs[0].runs:
                run.bold = True
    doc.add_heading("Bill to", level=1)
    for value in (_text(v, "recipient_name"), _text(v, "recipient_organization"), _text(v, "recipient_address")):
        if value:
            doc.add_paragraph(value)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ("Description", "Qty", f"Unit price ({currency})", f"Amount ({currency})")):
        cell.text = label
        _set_cell_shading(cell, "DCE8FA")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    subtotal = Decimal("0")
    for description, qty, unit_price in items:
        amount = qty * unit_price
        subtotal += amount
        row = table.add_row().cells
        row[0].text = description
        row[1].text = str(qty.normalize())
        row[2].text = _fmt_decimal(unit_price)
        row[3].text = _fmt_decimal(amount)
    total_row = table.add_row().cells
    total_row[0].merge(total_row[2])
    total_row[0].text = "TOTAL DUE"
    total_row[3].text = f"{currency} {_fmt_decimal(subtotal)}"
    for cell in total_row:
        _set_cell_shading(cell, "EAF1FB")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    _add_section(doc, "Payment terms", _text(v, "payment_terms"))
    _add_section(doc, "Payment details", _text(v, "payment_details"))
    notice = doc.add_paragraph("Please contact the sender promptly if you have any questions about this invoice.")
    for run in notice.runs:
        run.italic = True
        run.font.color.rgb = RGBColor(95, 112, 136)
        run.font.size = Pt(9)


def _build_proposal(doc: Document, v: dict[str, str]) -> None:
    _require(v, ["project_title", "recipient_name"], "For a proposal, enter the project title and recipient.")
    if not any(_text(v, key) for key in ("executive_summary", "objectives", "scope", "deliverables")):
        raise DocumentBuilderError("Add at least a summary, objectives, scope or deliverables for the proposal.")
    _top_identity(doc, v, _text(v, "project_title"))
    recipient = _text(v, "recipient_name")
    if _text(v, "recipient_organization"):
        recipient += "\n" + _text(v, "recipient_organization")
    doc.add_paragraph("Prepared for: " + recipient)
    _add_section(doc, "Executive summary", _text(v, "executive_summary"))
    _add_section(doc, "Objectives", _text(v, "objectives"), bullets=True)
    _add_section(doc, "Scope of work", _text(v, "scope"))
    _add_section(doc, "Deliverables", _text(v, "deliverables"), bullets=True)
    _add_section(doc, "Timeline and milestones", _text(v, "timeline"))
    _add_section(doc, "Investment and payment terms", _text(v, "investment"))
    _add_section(doc, "Assumptions and exclusions", _text(v, "assumptions"), bullets=True)
    _add_section(doc, "Proposed next steps", _text(v, "next_steps"), bullets=True)


def build_document(doc_type: str, values: dict[str, str]) -> tuple[bytes, str]:
    if doc_type not in DOCUMENT_TYPES:
        raise DocumentBuilderError("Choose a document template from the list.")
    for key, value in values.items():
        if len(value) > MAX_TEXT_FIELD:
            raise DocumentBuilderError("One of the fields is too long. Keep each field under 20,000 characters.")
    title = _text(values, "document_title") or DOCUMENT_TYPES[doc_type]
    if doc_type == "proposal":
        title = _text(values, "project_title", title)
    elif doc_type == "meeting-minutes":
        title = _text(values, "meeting_title", title)
    elif doc_type == "resume":
        title = f"{_text(values, 'full_name', 'Resume')} CV"
    elif doc_type == "invoice" and _text(values, "invoice_number"):
        title = f"Invoice {_text(values, 'invoice_number')}"
    author = _text(values, "full_name") or _text(values, "organization") or "Papyronix Document Builder"
    doc = _new_document(title, author, _text(values, "organization"), _text(values, "document_date"))
    builders = {
        "business-letter": _build_letter,
        "project-report": _build_report,
        "meeting-minutes": _build_minutes,
        "resume": _build_resume,
        "invoice": _build_invoice,
        "proposal": _build_proposal,
    }
    builders[doc_type](doc, values)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue(), title
