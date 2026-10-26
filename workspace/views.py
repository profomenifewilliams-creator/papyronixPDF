from __future__ import annotations

from pathlib import Path
from time import monotonic
import csv
import io
import re

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.db import connection
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import content_disposition_header
from django.utils import timezone
from django.views.decorators.http import require_POST, require_http_methods
from pypdf import PdfReader

from .forms import RegistrationForm
from .models import DocumentAsset, ProcessingJob
from .services import TOOL_MAP, TOOLS, ProcessingError, process_tool, outputs_as_zip, safe_stem

STOP_WORDS = set("a an the and or but if then than is are was were be been being to of for with in on at by from as it its this that these those you your we our they them their what when where why how can could should would will may might has have had do does did not no about into across between over under after before is are from which who whose also any all some each per use using find show list compare summarize summary explain document documents file files page pages pdf with without main key important information details details terms clause clauses".split())


def health(request):
    """Render health check without exposing configuration or database details."""
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            cursor.fetchone()
    except Exception:
        return JsonResponse({"status": "unavailable"}, status=503)
    return JsonResponse({"status": "ok"})
ALLOWED_DOC_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".pptx"}
MAX_LIBRARY_UPLOAD = 25 * 1024 * 1024


def _clean_filename(value: str) -> str:
    # Treat Windows and POSIX separators consistently, even when the server OS differs.
    name = str(value or "document").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip(" .")
    return name[:180] or "document"


def _format_size(value: int) -> str:
    if value < 1024:
        return f"{value} B"
    if value < 1024 * 1024:
        return f"{value / 1024:.1f} KB"
    return f"{value / (1024 * 1024):.2f} MB"


def home(request):
    recent = ProcessingJob.objects.filter(user=request.user)[:4] if request.user.is_authenticated else []
    counts = {
        "documents": DocumentAsset.objects.filter(user=request.user).count() if request.user.is_authenticated else 0,
        "jobs": ProcessingJob.objects.filter(user=request.user, succeeded=True).count() if request.user.is_authenticated else 0,
    }
    featured_slugs = ["merge", "pdf-to-word", "word-template-fill", "word-compare", "word-to-pdf", "excel-compare", "excel-to-word", "pdf-ocr"]
    featured = [TOOL_MAP[slug] for slug in featured_slugs if slug in TOOL_MAP]
    return render(request, "home.html", {"tools": featured, "recent": recent, "counts": counts})


def studio(request):
    groups = {}
    for tool in TOOLS:
        groups.setdefault(tool["group"], []).append(tool)
    return render(request, "studio.html", {"groups": groups, "tools": TOOLS})



@require_http_methods(["GET", "POST"])
def document_builder(request):
    from .document_builder import DOCUMENT_TYPES, DocumentBuilderError, build_document

    if request.method == "GET":
        return render(request, "builder.html", {
            "document_types": DOCUMENT_TYPES,
            "form_values": {"document_date": timezone.localdate().isoformat(), "currency": "NGN"},
            "error": "",
        })

    values = {key: value.strip() for key, value in request.POST.items() if key not in {"csrfmiddlewaretoken", "doc_type", "output_format"}}
    if sum(len(value) for value in values.values()) > 100_000:
        return render(request, "builder.html", {"document_types": DOCUMENT_TYPES, "form_values": request.POST, "error": "The combined form content is too long. Reduce the amount of text and try again."}, status=400)

    doc_type = request.POST.get("doc_type", "")
    field_aliases = {
        "business-letter": {"recipient_name": "recipient_name_letter", "recipient_organization": "recipient_organization_letter", "recipient_address": "recipient_address_letter"},
        "project-report": {"executive_summary": "report_executive_summary"},
        "meeting-minutes": {},
        "resume": {},
        "invoice": {"recipient_name": "recipient_name_invoice", "recipient_organization": "recipient_organization_invoice", "recipient_address": "recipient_address_invoice"},
        "proposal": {"recipient_name": "recipient_name_proposal", "recipient_organization": "recipient_organization_proposal", "executive_summary": "proposal_executive_summary"},
    }
    for canonical, posted_key in field_aliases.get(doc_type, {}).items():
        values[canonical] = (request.POST.get(posted_key) or "").strip()
    output_format = request.POST.get("output_format", "docx").lower()
    if output_format not in {"docx", "pdf"}:
        return render(request, "builder.html", {"document_types": DOCUMENT_TYPES, "form_values": request.POST, "error": "Choose Word or PDF as the output format."}, status=400)

    started = monotonic()
    operation = f"build:{doc_type}"[:40]
    try:
        docx_content, title = build_document(doc_type, values)
        if output_format == "pdf":
            from .office_services import _local_libreoffice_convert
            content = _local_libreoffice_convert(f"{safe_stem(title)}.docx", docx_content)
            content_type = "application/pdf"
        else:
            content = docx_content
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"{safe_stem(title)}.{output_format}"
        elapsed_ms = int((monotonic() - started) * 1000)
        _log_job(request, operation, [], 0, elapsed_ms, True, filename)
        response = HttpResponse(content, content_type=content_type)
        response["Content-Disposition"] = content_disposition_header(True, filename)
        response["X-Content-Type-Options"] = "nosniff"
        return response
    except (DocumentBuilderError, ProcessingError) as exc:
        elapsed_ms = int((monotonic() - started) * 1000)
        _log_job(request, operation, [], 0, elapsed_ms, False, error_code="validation")
        return render(request, "builder.html", {"document_types": DOCUMENT_TYPES, "form_values": request.POST, "error": str(exc)}, status=400)
    except Exception:
        elapsed_ms = int((monotonic() - started) * 1000)
        _log_job(request, operation, [], 0, elapsed_ms, False, error_code="processing")
        return render(request, "builder.html", {"document_types": DOCUMENT_TYPES, "form_values": request.POST, "error": "The document could not be generated. Check the fields and try again."}, status=422)

def tool_page(request, slug):
    tool = TOOL_MAP.get(slug)
    if not tool:
        raise Http404("Document tool not found")
    return render(request, "tools/tool.html", {"tool": tool, "form_values": {}, "error": ""})


def _log_job(request, operation, file_names, byte_total, elapsed_ms, succeeded, output_name="", error_code=""):
    # Only persist account-linked activity; uploaded payloads are not logged in job history.
    if request.user.is_authenticated:
        ProcessingJob.objects.create(
            user=request.user, operation=operation,
            input_names=[_clean_filename(name)[:255] for name in file_names[:30]],
            output_name=output_name[:255], input_bytes=byte_total,
            elapsed_ms=min(max(int(elapsed_ms), 0), 2_147_483_647),
            succeeded=succeeded, error_code=error_code[:40],
        )


@require_POST
def process_tool(request):
    operation = request.POST.get("operation", "").strip()
    tool = TOOL_MAP.get(operation)
    if not tool:
        return render(request, "tools/tool.html", {"tool": TOOLS[0], "form_values": request.POST, "error": "Choose a supported operation."}, status=400)
    uploads = request.FILES.getlist("files")
    if not uploads:
        return render(request, "tools/tool.html", {"tool": tool, "form_values": request.POST, "error": "Choose at least one file to begin."}, status=400)
    if not tool.get("multi", False) and len(uploads) > 1:
        return render(request, "tools/tool.html", {"tool": tool, "form_values": request.POST, "error": "This operation accepts one source file at a time."}, status=400)
    byte_total = sum(upload.size for upload in uploads)
    if byte_total > settings.PAPYRONIX_MAX_UPLOAD_BYTES:
        limit = _format_size(settings.PAPYRONIX_MAX_UPLOAD_BYTES)
        return render(request, "tools/tool.html", {"tool": tool, "form_values": request.POST, "error": f"The combined upload exceeds this local instance's {limit} limit."}, status=413)
    if any(upload.size == 0 for upload in uploads):
        return render(request, "tools/tool.html", {"tool": tool, "form_values": request.POST, "error": "One of the selected files is empty."}, status=400)
    started = monotonic()
    filenames = [_clean_filename(upload.name) for upload in uploads]
    try:
        files = [(_clean_filename(upload.name), upload.read()) for upload in uploads]
        params = {key: value for key, value in request.POST.items()}
        artifacts = process_tool(operation, files, params)
        elapsed_ms = int((monotonic() - started) * 1000)
        if len(artifacts) == 1:
            artifact = artifacts[0]
            response = HttpResponse(artifact.content, content_type=artifact.content_type)
            response["Content-Disposition"] = content_disposition_header(True, artifact.name)
            response["X-Content-Type-Options"] = "nosniff"
            _log_job(request, operation, filenames, byte_total, elapsed_ms, True, artifact.name)
            return response
        zipped = outputs_as_zip(artifacts)
        output_name = f"{safe_stem(filenames[0])}-outputs.zip"
        response = HttpResponse(zipped, content_type="application/zip")
        response["Content-Disposition"] = content_disposition_header(True, output_name)
        response["X-Content-Type-Options"] = "nosniff"
        _log_job(request, operation, filenames, byte_total, elapsed_ms, True, output_name)
        return response
    except ProcessingError as exc:
        elapsed_ms = int((monotonic() - started) * 1000)
        _log_job(request, operation, filenames, byte_total, elapsed_ms, False, error_code="validation")
        return render(request, "tools/tool.html", {"tool": tool, "form_values": request.POST, "error": str(exc)}, status=400)
    except Exception:
        elapsed_ms = int((monotonic() - started) * 1000)
        _log_job(request, operation, filenames, byte_total, elapsed_ms, False, error_code="processing")
        return render(request, "tools/tool.html", {"tool": tool, "form_values": request.POST, "error": "The file could not be processed. Check that it is valid and not damaged, then try again."}, status=422)


def _extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(content), strict=False)
        if reader.is_encrypted:
            raise ProcessingError("Password-protected PDFs cannot be indexed into the document library. Use Document Studio's Unlock tool first if you know the password.")
        if len(reader.pages) > 1000:
            raise ProcessingError("Document library indexing is limited to PDFs with 1,000 pages.")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            return "[No selectable text was found. This may be a scanned PDF; use Document Studio’s OCR scanned PDF tool if local Tesseract is installed.]"
        return text[:300_000]
    if extension == ".txt" or extension == ".md":
        return content.decode("utf-8", errors="replace")[:300_000]
    if extension == ".csv":
        decoded = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(decoded)))
        return "\n".join(" | ".join(cell for cell in row) for row in rows)[:300_000]
    if extension == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)[:300_000]
    if extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        parts = []
        try:
            for sheet in workbook.worksheets[:30]:
                parts.append(f"SHEET: {sheet.title}")
                for row in sheet.iter_rows(max_row=5000, max_col=100, values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        parts.append(" | ".join(values))
        finally:
            workbook.close()
        return "\n".join(parts)[:300_000]
    if extension == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(content))
        lines = []
        for index, slide in enumerate(presentation.slides, start=1):
            lines.append(f"SLIDE {index}")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    lines.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)[:300_000]
    raise ProcessingError("This file type is not supported by the document library.")


@login_required
def library(request):
    if request.method == "POST":
        upload = request.FILES.get("document")
        if not upload:
            messages.error(request, "Choose a file to upload.")
            return redirect("library")
        filename = _clean_filename(upload.name)
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_DOC_EXTENSIONS:
            messages.error(request, "Supported library formats: PDF, DOCX, TXT, MD, CSV, XLSX and PPTX.")
            return redirect("library")
        if upload.size <= 0 or upload.size > MAX_LIBRARY_UPLOAD:
            messages.error(request, "Documents must be smaller than 25 MB and cannot be empty.")
            return redirect("library")
        raw = upload.read()
        try:
            extracted = _extract_document_text(filename, raw)
        except ProcessingError as exc:
            messages.error(request, str(exc))
            return redirect("library")
        except Exception:
            messages.error(request, "The file could not be indexed. Check it is not damaged, then try again.")
            return redirect("library")
        upload.seek(0)
        asset = DocumentAsset(user=request.user, original_name=filename, content_type=upload.content_type or "application/octet-stream", size_bytes=upload.size, extracted_text=extracted)
        asset.file.save(filename, upload, save=True)
        messages.success(request, f"{filename} was added to your private document library.")
        return redirect("library")

    query = request.GET.get("q", "").strip()[:100]
    documents = DocumentAsset.objects.filter(user=request.user)
    if query:
        documents = documents.filter(original_name__icontains=query)
    return render(request, "workspace/library.html", {"documents": documents[:100], "query": query, "total_documents": DocumentAsset.objects.filter(user=request.user).count(), "max_upload_mb": 25})


@login_required
def download_document(request, pk):
    document = get_object_or_404(DocumentAsset, pk=pk, user=request.user)
    try:
        document.file.open("rb")
    except (FileNotFoundError, OSError):
        raise Http404("The file is no longer available on this server.")
    response = FileResponse(document.file, as_attachment=True, filename=_clean_filename(document.original_name), content_type=document.content_type or "application/octet-stream")
    response["X-Content-Type-Options"] = "nosniff"
    return response


@login_required
@require_POST
def delete_document(request, pk):
    document = get_object_or_404(DocumentAsset, pk=pk, user=request.user)
    document.file.delete(save=False)
    document.delete()
    messages.success(request, "Document removed from your private library.")
    return redirect("library")


def _local_answer(question: str, documents: list[DocumentAsset]) -> tuple[str, list[dict]]:
    terms = [word.lower() for word in re.findall(r"[A-Za-z0-9][A-Za-z0-9'-]{1,}", question) if word.lower() not in STOP_WORDS]
    terms = list(dict.fromkeys(terms))[:30]
    if not documents:
        return "Upload a document to build a local, searchable knowledge space. The assistant runs on this server and does not call a cloud AI API.", []
    if not terms:
        terms = [word.lower() for word in re.findall(r"[A-Za-z0-9][A-Za-z0-9'-]{2,}", question)][:10]
    snippets = []
    for doc in documents[:100]:
        text = (doc.extracted_text or "")[:60_000]
        sentences = [sentence.strip() for sentence in re.split(r"(?<=[.!?])\s+|\n+", text) if len(sentence.strip()) > 24]
        ranked = []
        for index, sentence in enumerate(sentences):
            lower = sentence.lower()
            matched = [term for term in terms if term in lower]
            score = sum(2 if re.search(rf"\b{re.escape(term)}\b", lower) else 1 for term in matched)
            if score:
                ranked.append((score, index, sentence, matched))
        for score, index, sentence, matched in sorted(ranked, key=lambda item: (-item[0], item[1]))[:3]:
            snippets.append({"document": doc.original_name, "text": sentence[:1100], "score": score, "matched": matched})
    snippets.sort(key=lambda item: item["score"], reverse=True)
    snippets = snippets[:10]
    if not snippets:
        answer = "I couldn't find a strong text match for that question in the selected documents. Try a few specific words from the topic, or check whether the file contains selectable text. This offline assistant retrieves relevant passages; it doesn't invent an answer when it has no matching evidence."
    else:
        answer = "I searched the selected documents locally and found these passages with the strongest word matches. Use them as source evidence; this offline version ranks and extracts text rather than generating new conclusions."
    return answer, snippets


@login_required
def assistant(request):
    documents = list(DocumentAsset.objects.filter(user=request.user)[:100])
    answer = ""
    snippets = []
    question = ""
    selected_ids = []
    if request.method == "GET" and request.GET.get("document", "").isdigit():
        requested_id = int(request.GET["document"])
        if any(document.pk == requested_id for document in documents):
            selected_ids = [requested_id]
    if request.method == "POST":
        question = request.POST.get("question", "").strip()[:1000]
        selected_ids = []
        for value in request.POST.getlist("document_ids")[:10]:
            if value.isdigit():
                item_id = int(value)
                if item_id not in selected_ids:
                    selected_ids.append(item_id)
            if len(selected_ids) >= 5:
                break
        selected = [document for document in documents if document.pk in selected_ids] if selected_ids else documents
        if question:
            answer, snippets = _local_answer(question, selected)
        else:
            messages.error(request, "Enter a question to search your documents.")
    return render(request, "workspace/assistant.html", {"documents": documents, "answer": answer, "snippets": snippets, "question": question, "selected_ids": selected_ids})


@login_required
def history(request):
    jobs = ProcessingJob.objects.filter(user=request.user)[:100]
    return render(request, "workspace/history.html", {"jobs": jobs})


def register(request):
    if request.user.is_authenticated:
        return redirect("studio")
    form = RegistrationForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        login(request, user)
        messages.success(request, "Your local Papyronix account is ready.")
        return redirect("studio")
    return render(request, "registration/register.html", {"form": form})
