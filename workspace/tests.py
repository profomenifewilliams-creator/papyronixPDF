from io import BytesIO
from unittest import TestCase

from PIL import Image
from pypdf import PdfReader, PdfWriter

from .services import ProcessingError, parse_pages, process_tool


def sample_pdf(pages=4):
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class PageRangeTests(TestCase):
    def test_all_pages(self):
        self.assertEqual(parse_pages("all", 4), [1, 2, 3, 4])

    def test_range_and_list(self):
        self.assertEqual(parse_pages("1,3-4", 4), [1, 3, 4])

    def test_duplicate_rejected(self):
        with self.assertRaises(ProcessingError):
            parse_pages("1,1", 4)

    def test_out_of_range_rejected(self):
        with self.assertRaises(ProcessingError):
            parse_pages("1-5", 4)


class LocalPdfEngineTests(TestCase):
    def setUp(self):
        self.pdf = sample_pdf(4)
        self.file = [("sample.pdf", self.pdf)]

    def test_extract_pages(self):
        result = process_tool("extract", self.file, {"pages": "1,3"})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 2)

    def test_split_ranges(self):
        result = process_tool("split", self.file, {"split_mode": "ranges", "ranges": "1-2,3-4"})
        self.assertEqual(len(result), 2)
        self.assertEqual([len(PdfReader(BytesIO(item.content)).pages) for item in result], [2, 2])

    def test_split_each_page(self):
        result = process_tool("split", self.file, {"split_mode": "each"})
        self.assertEqual(len(result), 4)
        self.assertTrue(all(len(PdfReader(BytesIO(item.content)).pages) == 1 for item in result))

    def test_remove_pages(self):
        result = process_tool("remove", self.file, {"pages": "2,4"})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 2)

    def test_reorder_pages(self):
        result = process_tool("reorder", self.file, {"order": "4,3,2,1"})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 4)

    def test_rotate(self):
        result = process_tool("rotate", self.file, {"pages": "1", "angle": "90"})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 4)

    def test_watermark(self):
        result = process_tool("watermark", self.file, {"watermark": "CONFIDENTIAL"})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 4)

    def test_page_numbers(self):
        result = process_tool("page-numbers", self.file, {"start_number": "10"})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 4)

    def test_optimize(self):
        result = process_tool("optimize", self.file, {})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 4)

    def test_merge(self):
        result = process_tool("merge", self.file + [("second.pdf", sample_pdf(2))], {})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 6)

    def test_password_protect_and_unlock(self):
        protected = process_tool("protect", self.file, {"new_password": "strong-test-pass", "confirm_password": "strong-test-pass"})[0]
        reader = PdfReader(BytesIO(protected.content))
        self.assertTrue(reader.is_encrypted)
        unlocked = process_tool("unlock", [("protected.pdf", protected.content)], {"pdf_password": "strong-test-pass"})[0]
        self.assertFalse(PdfReader(BytesIO(unlocked.content)).is_encrypted)

    def test_invalid_pdf_is_rejected(self):
        with self.assertRaises(ProcessingError):
            process_tool("extract", [("bad.pdf", b"not a pdf")], {"pages": "1"})

    def test_images_to_pdf(self):
        image = Image.new("RGB", (80, 60), color=(230, 240, 255))
        stream = BytesIO(); image.save(stream, format="PNG")
        result = process_tool("images-to-pdf", [("sample.png", stream.getvalue())], {})
        self.assertEqual(len(PdfReader(BytesIO(result[0].content)).pages), 1)


class LocalDocumentEngineTests(TestCase):
    """Pure local-format conversion tests; do not require any AI/network credentials."""
    def make_text_pdf(self, text="Papyronix local document test"):
        from reportlab.pdfgen import canvas
        output = BytesIO()
        page = canvas.Canvas(output)
        page.drawString(72, 720, text)
        page.showPage()
        page.save()
        return output.getvalue()

    def make_docx(self, paragraphs=("A sample document", "Target phrase")):
        from docx import Document
        output = Document()
        for line in paragraphs:
            output.add_paragraph(line)
        stream = BytesIO()
        output.save(stream)
        return stream.getvalue()

    def make_xlsx(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["Name", "Count"])
        sheet.append(["Sample", 1])
        sheet.append(["Sample", 1])
        sheet.append([None, None])
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def make_pptx(self):
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Sample presentation heading"
        stream = BytesIO()
        presentation.save(stream)
        return stream.getvalue()

    def test_pdf_to_text_extracts_selectable_text(self):
        result = process_tool("pdf-to-text", [("source.pdf", self.make_text_pdf())], {})[0]
        self.assertIn(b"Papyronix local document test", result.content)
        self.assertEqual(result.content_type, "text/plain; charset=utf-8")

    def test_pdf_to_word_creates_docx(self):
        from docx import Document
        result = process_tool("pdf-to-word", [("source.pdf", self.make_text_pdf())], {})[0]
        exported = Document(BytesIO(result.content))
        self.assertIn("Papyronix local document test", " ".join(p.text for p in exported.paragraphs))

    def test_pdf_to_images_returns_png(self):
        result = process_tool("pdf-to-images", [("source.pdf", self.make_text_pdf())], {"dpi": "72", "pages": "all"})[0]
        self.assertTrue(result.content.startswith(b"\x89PNG\r\n\x1a\n"))
        self.assertEqual(result.content_type, "image/png")

    def test_word_find_replace(self):
        from docx import Document
        source = self.make_docx(("Alpha phrase",))
        result = process_tool("word-replace", [("source.docx", source)], {"find_text": "Alpha", "replace_text": "Beta"})[0]
        exported = Document(BytesIO(result.content))
        self.assertEqual(exported.paragraphs[0].text, "Beta phrase")

    def test_word_replace_same_text_terminates(self):
        source = self.make_docx(("Alpha",))
        result = process_tool("word-replace", [("source.docx", source)], {"find_text": "Alpha", "replace_text": "Alpha"})[0]
        self.assertTrue(result.name.endswith(".docx"))

    def test_text_to_word(self):
        from docx import Document
        result = process_tool("text-to-word", [("source.txt", b"First paragraph\n\nSecond paragraph")], {})[0]
        exported = Document(BytesIO(result.content))
        self.assertEqual([p.text for p in exported.paragraphs], ["First paragraph", "Second paragraph"])

    def test_word_count_report(self):
        result = process_tool("word-count", [("source.docx", self.make_docx(("one two three",)))], {})[0]
        self.assertIn(b"Words: 3", result.content)

    def test_csv_to_excel_preserves_leading_zero_and_formula_text(self):
        from openpyxl import load_workbook
        csv_data = b"id,value\n001,2\n=1+1,3\n"
        result = process_tool("csv-to-excel", [("source.csv", csv_data)], {})[0]
        workbook = load_workbook(BytesIO(result.content))
        self.assertEqual(workbook.active["A2"].value, "001")
        self.assertEqual(workbook.active["A3"].data_type, "s")
        workbook.close()

    def test_excel_to_csv_exports_sheet_values(self):
        result = process_tool("excel-to-csv", [("source.xlsx", self.make_xlsx())], {})
        self.assertEqual(len(result), 1)
        self.assertIn(b"Sample,1", result[0].content)

    def test_excel_clean_removes_duplicate_and_blank_rows(self):
        from openpyxl import load_workbook
        result = process_tool("excel-clean", [("source.xlsx", self.make_xlsx())], {})[0]
        workbook = load_workbook(BytesIO(result.content), read_only=True)
        rows = list(workbook.active.values)
        workbook.close()
        self.assertEqual(rows, [("Name", "Count"), ("Sample", 1)])

    def test_excel_merge_combines_worksheets(self):
        from openpyxl import load_workbook
        source = self.make_xlsx()
        result = process_tool("excel-merge", [("first.xlsx", source), ("second.xlsx", source)], {})[0]
        workbook = load_workbook(BytesIO(result.content), read_only=True)
        self.assertEqual(len(workbook.worksheets), 2)
        workbook.close()

    def test_pptx_to_text(self):
        result = process_tool("pptx-to-text", [("source.pptx", self.make_pptx())], {})[0]
        self.assertIn(b"Sample presentation heading", result.content)

    def test_multi_artifact_zip(self):
        from .services import outputs_as_zip
        source = self.make_text_pdf()
        artifacts = process_tool("pdf-to-images", [("source.pdf", source)], {"dpi": "72", "pages": "all"})
        zipped = outputs_as_zip(artifacts)
        self.assertTrue(zipped.startswith(b"PK"))


class LocalDocumentBuilderTests(TestCase):
    def test_business_letter_template_builds_docx(self):
        from docx import Document
        from .document_builder import build_document
        content, title = build_document("business-letter", {
            "recipient_name": "Jordan Client", "subject": "Service update",
            "body": "The project is progressing as agreed.", "full_name": "Alex Author",
        })
        document = Document(BytesIO(content))
        text = " ".join(p.text for p in document.paragraphs)
        self.assertIn("Jordan Client", text)
        self.assertIn("Service update", text)
        self.assertIn("The project is progressing", text)
        self.assertEqual(title, "Business letter")

    def test_invoice_template_calculates_line_amounts(self):
        from docx import Document
        from .document_builder import build_document
        content, title = build_document("invoice", {
            "full_name": "Local Vendor", "recipient_name": "Client",
            "invoice_number": "INV-42", "currency": "NGN",
            "line_items": "Design work | 2 | 1500\nSupport | 1 | 500",
        })
        document = Document(BytesIO(content))
        table_text = " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("3,000.00", table_text)
        self.assertIn("3,500.00", table_text)
        self.assertEqual(title, "Invoice INV-42")

    def test_resume_template_builds_structured_sections(self):
        from docx import Document
        from .document_builder import build_document
        content, _ = build_document("resume", {
            "full_name": "Alex Candidate", "skills": "Python\nDjango",
            "experience": "Software developer", "education": "BSc Computing",
        })
        document = Document(BytesIO(content))
        text = " ".join(p.text for p in document.paragraphs)
        self.assertIn("Alex Candidate", text)
        self.assertIn("Key skills", text)
        self.assertIn("Work experience", text)

    def test_invalid_template_input_is_rejected(self):
        from .document_builder import DocumentBuilderError, build_document
        with self.assertRaises(DocumentBuilderError):
            build_document("invoice", {"full_name": "Vendor", "recipient_name": "Client", "invoice_number": "1", "line_items": "not a valid row"})

class IndependentOfficeExpansionTests(TestCase):
    """Regression tests for the local-only Office/document expansion."""
    def make_docx(self, paragraphs=("A sample",)):
        from docx import Document
        doc = Document()
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        stream = BytesIO()
        doc.save(stream)
        return stream.getvalue()

    def make_workbook(self, rows=(('Name', 'Score'), ('Mia', 7), ('Ayo', 3))):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        for row in rows:
            sheet.append(list(row))
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def test_word_mail_merge_from_csv_generates_per_row_documents(self):
        from docx import Document
        template = self.make_docx(("Dear {{FullName}},", "Your unit is {{Department}}."))
        rows = b"FullName,Department\nAda Okafor,Operations\nMusa Bello,Finance\n"
        outputs = process_tool("word-template-fill", [("contacts.csv", rows), ("template.docx", template)], {})
        self.assertEqual(len(outputs), 2)
        text1 = " ".join(p.text for p in Document(BytesIO(outputs[0].content)).paragraphs)
        text2 = " ".join(p.text for p in Document(BytesIO(outputs[1].content)).paragraphs)
        self.assertIn("Ada Okafor", text1)
        self.assertIn("Musa Bello", text2)
        self.assertNotIn("{{FullName}}", text1)

    def test_word_mail_merge_requires_matching_csv_header(self):
        with self.assertRaises(ProcessingError):
            process_tool("word-template-fill", [("template.docx", self.make_docx(("Hello {{Client}}",))), ("data.csv", b"Person\nAyo\n")], {})

    def test_word_to_markdown_preserves_heading_emphasis_and_table(self):
        from docx import Document
        document = Document()
        document.add_heading("Project Summary", level=1)
        paragraph = document.add_paragraph()
        paragraph.add_run("Important text").bold = True
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Role"
        table.cell(1, 0).text = "Ayo"
        table.cell(1, 1).text = "Reviewer"
        stream = BytesIO(); document.save(stream)
        result = process_tool("word-to-markdown", [("summary.docx", stream.getvalue())], {})[0]
        self.assertIn(b"# Project Summary", result.content)
        self.assertIn(b"**Important text**", result.content)
        self.assertIn(b"| Ayo | Reviewer |", result.content)

    def test_word_extract_images(self):
        from docx import Document
        image = Image.new("RGB", (30, 20), color=(240, 240, 240))
        image_bytes = BytesIO(); image.save(image_bytes, format="PNG")
        document = Document(); document.add_picture(BytesIO(image_bytes.getvalue()))
        stream = BytesIO(); document.save(stream)
        results = process_tool("word-extract-images", [("illustrated.docx", stream.getvalue())], {})
        self.assertEqual(len(results), 1)
        self.assertTrue(results[0].name.lower().endswith(".png"))
        self.assertTrue(results[0].content.startswith(b"\x89PNG\r\n\x1a\n"))

    def test_markdown_to_word_formats_headings_lists_and_tables(self):
        from docx import Document
        source = b"# Report\n\nThis has **bold** text.\n\n- First\n- Second\n\n| Name | Score |\n| --- | --- |\n| Ayo | 8 |"
        result = process_tool("text-to-word", [("report.md", source)], {})[0]
        document = Document(BytesIO(result.content))
        self.assertEqual(document.paragraphs[0].style.name, "Heading 1")
        self.assertTrue(any(run.bold for paragraph in document.paragraphs for run in paragraph.runs))
        self.assertTrue(any("List Bullet" in paragraph.style.name for paragraph in document.paragraphs))
        self.assertEqual(len(document.tables), 1)

    def test_word_compare_reports_change(self):
        result = process_tool("word-compare", [("old.docx", self.make_docx(("Original phrase",))), ("new.docx", self.make_docx(("Revised phrase",)))], {})[0]
        self.assertIn(b"Original phrase", result.content)
        self.assertIn(b"Revised phrase", result.content)

    def test_word_table_to_excel_exports_tables(self):
        from docx import Document
        from openpyxl import load_workbook
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Code"
        table.cell(0, 1).text = "Name"
        table.cell(1, 0).text = "001"
        table.cell(1, 1).text = "Example"
        stream = BytesIO(); doc.save(stream)
        result = process_tool("word-table-to-excel", [("table.docx", stream.getvalue())], {})[0]
        workbook = load_workbook(BytesIO(result.content), data_only=False)
        self.assertEqual(workbook.active["A2"].value, "001")
        workbook.close()

    def test_word_page_number_field_added(self):
        from docx import Document
        result = process_tool("word-add-page-numbers", [("source.docx", self.make_docx())], {})[0]
        exported = Document(BytesIO(result.content))
        self.assertIn("PAGE", exported.sections[0].footer._element.xml)

    def test_word_clear_metadata(self):
        from docx import Document
        original = Document(BytesIO(self.make_docx()))
        original.core_properties.author = "Private Name"
        source = BytesIO(); original.save(source)
        result = process_tool("word-clear-metadata", [("source.docx", source.getvalue())], {})[0]
        cleared = Document(BytesIO(result.content))
        self.assertEqual(cleared.core_properties.author, "")

    def test_text_to_pdf_is_readable(self):
        from pypdf import PdfReader
        result = process_tool("text-to-pdf", [("guide.txt", b"# Local Guide\n\nThis is a local PDF output.\n- First item\n- Second item")], {})[0]
        reader = PdfReader(BytesIO(result.content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        self.assertIn("Local Guide", text)
        self.assertIn("local PDF output", text)

    def test_excel_split_sheets_returns_each_sheet(self):
        from openpyxl import Workbook, load_workbook
        wb = Workbook()
        wb.active.title = "First"
        wb.active.append(["Id"]); wb.active.append(["001"])
        second = wb.create_sheet("Second")
        second.append(["Flag"]); second.append(["Yes"])
        stream = BytesIO(); wb.save(stream); wb.close()
        outputs = process_tool("excel-split-sheets", [("book.xlsx", stream.getvalue())], {})
        self.assertEqual(len(outputs), 2)
        self.assertTrue(all(output.content.startswith(b"PK") for output in outputs))
        self.assertEqual(load_workbook(BytesIO(outputs[1].content), read_only=True).active.title, "Second")

    def test_excel_sort_by_header(self):
        from openpyxl import load_workbook
        source = self.make_workbook((("Name", "Score"), ("Mia", 7), ("Ayo", 3), ("Zee", 5)))
        result = process_tool("excel-sort", [("scores.xlsx", source)], {"sort_column": "Score", "sort_direction": "ascending"})[0]
        wb = load_workbook(BytesIO(result.content), data_only=False)
        values = [wb.active.cell(row, 2).value for row in range(2, 5)]
        wb.close()
        self.assertEqual(values, [3, 5, 7])

    def test_excel_compare_reports_changed_cells(self):
        old = self.make_workbook((("Name", "Score"), ("Mia", 7)))
        new = self.make_workbook((("Name", "Score"), ("Mia", 9)))
        result = process_tool("excel-compare", [("old.xlsx", old), ("new.xlsx", new)], {})[0]
        self.assertIn(b"B2", result.content)
        self.assertIn(b"Changed", result.content)

    def test_excel_to_word_creates_report(self):
        from docx import Document
        result = process_tool("excel-to-word", [("source.xlsx", self.make_workbook())], {})[0]
        doc = Document(BytesIO(result.content))
        self.assertTrue(any(table.rows[1].cells[0].text == "Mia" for table in doc.tables))

    def test_csv_clean_trims_and_deduplicates(self):
        result = process_tool("csv-clean", [("source.csv", b"Name,City\n Ayo ,Lagos\nAyo,Lagos\n,,\nMia,Abuja\n")], {})[0]
        self.assertEqual(result.content.decode("utf-8-sig").splitlines(), ["Name,City", "Ayo,Lagos", "Mia,Abuja"])

    def test_csv_filter_contains(self):
        result = process_tool("csv-filter", [("source.csv", b"Name,Status\nAyo,Approved\nMia,Pending\nTolu,Approved - Review\n")], {"filter_column": "Status", "filter_value": "approved", "filter_mode": "contains"})[0]
        lines = result.content.decode("utf-8-sig").splitlines()
        self.assertEqual(len(lines), 3)
        self.assertIn("Tolu,Approved - Review", lines)

    def test_csv_clean_neutralizes_spreadsheet_formula_text(self):
        result = process_tool("csv-clean", [("source.csv", b"Name,Value\nAyo,=1+1\nMia,-42\n")], {})[0]
        text = result.content.decode("utf-8-sig")
        self.assertIn("'=1+1", text)
        self.assertIn("-42", text)
